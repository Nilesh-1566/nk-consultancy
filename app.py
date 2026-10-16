from flask import Flask, render_template, request, send_file, jsonify, redirect, url_for
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import io

app = Flask(__name__)

# १. मुख्य होमपेज
@app.route('/')
def home():
    return render_template('home.html')

# २. दस्त मसुदा मेनू
@app.route('/draft')
def draft():
    return render_template('draft_menu.html')

# ३. खरेदीखत उप-मेनू
@app.route('/kharedi_menu')
def kharedi_menu():
    return render_template('kharedi_menu.html')

# ४. मूळ नावाने थेट फॉर्म उघडणारे डायनॅमिक Route
@app.route('/draft_form')
def draft_form():
    doc_type = request.args.get('type', 'प्लॉट खरेदीपत्र')
    return render_template('draft_form.html', doc_title=doc_type)

@app.route('/plot_draft')
def plot_draft():
    return render_template('draft_form.html', doc_title='प्लॉट खरेदीपत्र')

@app.route('/karar_to_kharedi_draft')
def karar_to_kharedi_draft():
    return render_template('draft_form.html', doc_title='करारपत्र वरून खरेदीपत्र')

# ५. Word (.docx) दस्त थेट तयार करून डाऊनलोड देणे
@app.route('/generate_custom_docx', methods=['POST'])
def generate_custom_docx():
    f = request.form
    doc_title = f.get('doc_title', 'खरेदीपत्र')
    doc = Document()

    # पक्षकार वर्गीकरण
    roles = f.getlist('party_role[]')
    names = f.getlist('party_name[]')
    ages = f.getlist('party_age[]')
    occs = f.getlist('party_occ[]')
    addrs = f.getlist('party_address[]')
    aadhars = f.getlist('party_aadhar[]')
    pans = f.getlist('party_pan[]')
    mobs = f.getlist('party_mobile[]')

    ghenar_list, denar_list, poa_list, manyata_list, jamin_list, sakshi_list = [], [], [], [], [], []

    for i in range(len(names)):
        item = {
            "name": names[i],
            "age": ages[i] if i < len(ages) else "",
            "occ": occs[i] if i < len(occs) else "",
            "addr": addrs[i] if i < len(addrs) else "",
            "aadhar": aadhars[i] if i < len(aadhars) else "",
            "pan": pans[i] if i < len(pans) else "",
            "mob": mobs[i] if i < len(mobs) else ""
        }
        role = roles[i] if i < len(roles) else "लिहून देणार"

        if role == 'लिहून घेणार':
            ghenar_list.append(item)
        elif role == 'लिहून देणार':
            denar_list.append(item)
        elif 'वटमुखत्यार' in role:
            poa_list.append(item)
        elif 'मान्यता' in role or 'संमती' in role:
            manyata_list.append(item)
        elif 'जामीनदार' in role:
            jamin_list.append(item)
        elif 'साक्षीदार' in role:
            sakshi_list.append(item)

    # -------------------------------------------------------------
    # १. तारण गहाणखत मसुदा
    # -------------------------------------------------------------
    if 'तारण' in doc_title or 'गहाण' in doc_title:
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_head.add_run("।। श्री ।।\n।। तारण गहाण खत ।।\n").bold = True

        doc.add_paragraph(
            f"तारण गहाण खत {f.get('village_taluka')} येथील मिळकतीचे रक्कम रुपये {f.get('amount_num')} "
            f"({f.get('amount_words')}) चे दि. {f.get('doc_date')}\n\n"
            f"{f.get('selected_bank')},\n"
            f"तर्फे {f.get('bank_rep_name')}\t\t\t\t\t\t\t\tलिहून घेणार\n"
        )

        p_d = doc.add_paragraph()
        p_d.add_run("कर्जदार (लिहून देणार):\n").bold = True
        for p in denar_list:
            p_d.add_run(f"{p['name']}, {p['age']}, धंदा- {p['occ']},\nरा. {p['addr']}\nआ.नं. {p['aadhar']} | पॅन: {p['pan']}\n\n")

        if jamin_list:
            p_j = doc.add_paragraph()
            p_j.add_run("जामीनदार:\n").bold = True
            for idx, p in enumerate(jamin_list):
                p_j.add_run(f"{idx+1}) {p['name']}, {p['age']}, धंदा- {p['occ']},\nरा. {p['addr']}\nआ.नं. {p['aadhar']}\n\n")

        doc.add_paragraph(
            f"कारणे लिहून देतो सतत तारण गहाण खत ऐसाजे कि,\n"
            f"१) मी कर्जदार {f.get('selected_bank')} या संस्थेचा सभासद असून {f.get('bank_meeting_res')} अन्वये "
            f"माझे नांवे रक्कम रुपये {f.get('amount_num')} ({f.get('amount_words')}) चे तारण कर्ज मंजूर झाले आहे.\n"
            f"२) सदर कर्जास व्याजाचा दर द.सा.द.शे. {f.get('bank_interest')} राहील व कर्जाची मुदत {f.get('bank_tenure')} राहील.\n\n"
            f"तारण मिळकतीचे वर्णन :-\n"
            f"{f.get('village_taluka')} येथील गट नं./सि.स.नं. {f.get('gat_no')} एकूण क्षेत्र {f.get('total_area')} पैकी "
            f"तारण दिलेले क्षेत्र {f.get('sold_area')} ची मिळकत सर्व.\n"
            f"चतुःसीमा :-\nपूर्वेस : {f.get('east')} | पश्चिमेस : {f.get('west')} | दक्षिणेस : {f.get('south')} | उत्तरेस : {f.get('north')}\n\n"
            f"येणेप्रमाणे तारण गहाणखत स्वखुषीने लिहून दिले असे सही दि. {f.get('doc_date')}."
        )

    # -------------------------------------------------------------
    # २. करारपत्र वरून खरेदीपत्र
    # -------------------------------------------------------------
    elif 'करार' in doc_title:
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_head.add_run("|| श्री स्वामी समर्थ ||\n|| खरेदीपत्र ||\n").bold = True

        doc.add_paragraph(
            f"रजि. करारपत्र {f.get('karar_ref')} रोजीच्या दस्तावरून खरेदीपत्र {f.get('village_taluka')} येथील मिळकतीचे "
            f"सरकारी बाजारभावाप्रमाणे रक्कम रुपये {f.get('amount_num')} ({f.get('amount_words')}) दि. {f.get('doc_date')} इसवीचे.\n\n"
        )

        p_g = doc.add_paragraph()
        p_g.add_run("खरेदी घेणार:\n").bold = True
        for p in ghenar_list:
            p_g.add_run(f"{p['name']}, {p['age']}, धंदा- {p['occ']},\nरा. {p['addr']}\nआ.नं. {p['aadhar']} | पॅन: {p['pan']}\n\n")

        p_d = doc.add_paragraph()
        p_d.add_run("खरेदी देणार:\n").bold = True
        for p in denar_list:
            p_d.add_run(f"{p['name']}, {p['age']}, धंदा- {p['occ']},\nरा. {p['addr']}\nआ.नं. {p['aadhar']}\n")
        if poa_list:
            p_d.add_run(f"यांचे तर्फे वटमुखत्यार म्हणून {poa_list[0]['name']}, {poa_list[0]['age']}\nरा. {poa_list[0]['addr']}\n\n")

        doc.add_paragraph(
            f"मिळकतीचे वर्णन :-\n"
            f"{f.get('village_taluka')} येथील गट नं. {f.get('gat_no')} एकूण क्षेत्र {f.get('total_area')} यापैकी "
            f"क्षेत्र {f.get('sold_area')} ची मिळकत सर्व.\n"
            f"चतुःसीमा :-\nपूर्वेस : {f.get('east')} | पश्चिमेस : {f.get('west')} | दक्षिणेस : {f.get('south')} | उत्तरेस : {f.get('north')}\n\n"
            f"येणेप्रमाणे करारपत्रावरून खरेदीपत्र लिहून दिले असे सही दि. {f.get('doc_date')}."
        )

    # -------------------------------------------------------------
    # ३. इतर सर्व दस्त (प्लॉट, शेती, सीटीएस खरेदीखत, हक्कसोड, बक्षीसपत्र, इ.)
    # -------------------------------------------------------------
    else:
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_head.add_run(f"|| श्री ||\n|| {doc_title} ||\n").bold = True

        doc.add_paragraph(
            f"{doc_title} {f.get('village_taluka')} येथील मिळकतीबाबतचे रक्कम रुपये {f.get('amount_num')} "
            f"({f.get('amount_words')}) दि. {f.get('doc_date')} इसवीचे.\n\n"
        )

        p_g = doc.add_paragraph()
        p_g.add_run("लिहून घेणार:\n").bold = True
        for p in ghenar_list:
            p_g.add_run(f"{p['name']}, {p['age']}, धंदा- {p['occ']},\nरा. {p['addr']}\nआ.नं. {p['aadhar']} | पॅन: {p['pan']}\n\n")

        p_d = doc.add_paragraph()
        p_d.add_run("लिहून देणार:\n").bold = True
        for p in denar_list:
            p_d.add_run(f"{p['name']}, {p['age']}, धंदा- {p['occ']},\nरा. {p['addr']}\nआ.नं. {p['aadhar']}\n")
        if poa_list:
            p_d.add_run(f"यांचे तर्फे वटमुखत्यार म्हणून {poa_list[0]['name']}, {poa_list[0]['age']}\nरा. {poa_list[0]['addr']}\n\n")

        if manyata_list:
            p_m = doc.add_paragraph()
            p_m.add_run("संमती / मान्यता देणार:\n").bold = True
            for p in manyata_list:
                p_m.add_run(f"{p['name']}, {p['age']}, रा. {p['addr']}\n")

        doc.add_paragraph(
            f"\nमिळकतीचे वर्णन :-\n"
            f"{f.get('village_taluka')} येथील मिळकत गट नं./सि.स.नं. {f.get('gat_no')} एकूण क्षेत्र {f.get('total_area')} पैकी "
            f"हस्तांतरित क्षेत्र {f.get('sold_area')} ची मिळकत सर्व.\n"
            f"चतुःसीमा :-\nपूर्वेस : {f.get('east')} | पश्चिमेस : {f.get('west')} | दक्षिणेस : {f.get('south')} | उत्तरेस : {f.get('north')}\n\n"
            f"येणेप्रमाणे {doc_title} स्वखुषीने लिहून दिले असे सही दि. {f.get('doc_date')}."
        )

    # Word फाईल मेमरीमध्ये साठवून डाऊनलोड पाठवणे
    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)

    clean_title = doc_title.replace(' ', '_')
    gat_clean = str(f.get('gat_no', '1')).replace('/', '_')
    filename = f"{clean_title}_Gat_{gat_clean}.docx"

    return send_file(
        target_stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

# ६. IGR व संपर्क Routes
@app.route('/igr')
def igr():
    return render_template('igr.html')

@app.route('/contact')
def contact():
    return render_template('contact.html')

if __name__ == '__main__':
    app.run(debug=True)
